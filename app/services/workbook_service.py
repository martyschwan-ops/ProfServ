"""
Workbook service – all reads and writes to the Excel .xlsx file.
The workbook is the system of record; no database is used.
"""
from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models.expense import Currency, Expense, ExpenseFilter, ExpenseType, PerDiemType, SourceType
from app.models.trip import Trip
from config import settings

logger = logging.getLogger(__name__)

# ── Column definitions ────────────────────────────────────────────────────────

TRIP_COLUMNS = [
    "Trip ID", "Trip Name", "Start Date", "End Date",
    "Notes", "Created At", "Updated At",
]

EXPENSE_COLUMNS = [
    "Expense ID", "Trip ID", "Trip Name", "Expense Date",
    "Vendor Name", "Expense Type", "Amount Original", "Currency Original",
    "Exchange Rate to CAD", "Amount CAD", "GST Original", "GST CAD",
    "Source Type", "Per Diem Type", "Receipt File Name", "Receipt File Path",
    "Notes", "Submitted Status", "Created At", "Updated At",
]

SETTINGS_COLUMNS = ["Key", "Value"]

HEADER_FILL = PatternFill("solid", fgColor="1F497D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _safe_date(val: Any) -> Optional[date]:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    try:
        from dateutil.parser import parse
        return parse(str(val)).date()
    except Exception:
        return None


def _safe_datetime(val: Any) -> datetime:
    if isinstance(val, datetime):
        return val
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    try:
        from dateutil.parser import parse
        return parse(str(val))
    except Exception:
        return datetime.utcnow()


def _safe_float(val: Any) -> Optional[float]:
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _safe_bool(val: Any) -> bool:
    if isinstance(val, bool):
        return val
    if isinstance(val, int):
        return bool(val)
    if isinstance(val, str):
        return val.lower() in ("true", "1", "yes")
    return False


def _style_header_row(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _set_column_widths(ws, widths: Dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


# ── WorkbookService ───────────────────────────────────────────────────────────

class WorkbookService:
    def __init__(self) -> None:
        self.path: Path = settings.workbook_path

    # ── Workbook lifecycle ────────────────────────────────────────────────────

    def bootstrap(self) -> None:
        """Create the workbook and all sheets if they don't exist."""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            wb = self._load()
            changed = False
            if "Trips" not in wb.sheetnames:
                self._init_trips_sheet(wb)
                changed = True
            if "Expenses" not in wb.sheetnames:
                self._init_expenses_sheet(wb)
                changed = True
            if "Settings" not in wb.sheetnames:
                self._init_settings_sheet(wb)
                changed = True
            if changed:
                self._save(wb)
        else:
            wb = Workbook()
            # Remove default sheet
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
            self._init_trips_sheet(wb)
            self._init_expenses_sheet(wb)
            self._init_settings_sheet(wb)
            self._save(wb)
            logger.info("Workbook created at %s", self.path)

    def _load(self) -> Workbook:
        try:
            return openpyxl.load_workbook(self.path)
        except Exception as exc:
            raise IOError(
                f"Cannot open workbook at {self.path}. "
                "Is it open in Excel? Close it and try again."
            ) from exc

    def _save(self, wb: Workbook) -> None:
        try:
            wb.save(self.path)
        except PermissionError as exc:
            raise IOError(
                f"Cannot save workbook at {self.path}. "
                "Is it open in Excel? Close it and try again."
            ) from exc

    # ── Sheet initialisation ──────────────────────────────────────────────────

    def _init_trips_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Trips")
        ws.append(TRIP_COLUMNS)
        _style_header_row(ws)
        _set_column_widths(ws, {1: 38, 2: 30, 3: 14, 4: 14, 5: 40, 6: 22, 7: 22})
        ws.freeze_panes = "A2"

    def _init_expenses_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Expenses")
        ws.append(EXPENSE_COLUMNS)
        _style_header_row(ws)
        ws.freeze_panes = "A2"
        widths = {
            1: 38, 2: 38, 3: 25, 4: 14, 5: 28, 6: 14,
            7: 16, 8: 12, 9: 20, 10: 14, 11: 14, 12: 10,
            13: 14, 14: 14, 15: 40, 16: 60, 17: 40, 18: 12,
            19: 22, 20: 22,
        }
        _set_column_widths(ws, widths)

    def _init_settings_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Settings")
        ws.append(SETTINGS_COLUMNS)
        _style_header_row(ws)
        defaults = [
            ("Breakfast Per Diem", settings.breakfast_amount),
            ("Lunch Per Diem", settings.lunch_amount),
            ("Dinner Per Diem", settings.dinner_amount),
            ("Base Currency", settings.base_currency),
        ]
        for row in defaults:
            ws.append(row)
        _set_column_widths(ws, {1: 30, 2: 20})

    # ── Sheet access guard ────────────────────────────────────────────────────

    def _get_sheet(self, wb: Workbook, name: str, init_fn):
        if name not in wb.sheetnames:
            logger.warning("Sheet '%s' missing – recreating", name)
            init_fn(wb)
        return wb[name]

    # ── Column index lookup ───────────────────────────────────────────────────

    @staticmethod
    def _col_map(ws) -> Dict[str, int]:
        """Return {header_name: column_index (1-based)} from row 1."""
        return {
            cell.value: cell.column
            for cell in ws[1]
            if cell.value is not None
        }

    # ═══════════════════════════════════════════════════════════════════════════
    # TRIPS
    # ═══════════════════════════════════════════════════════════════════════════

    def list_trips(self) -> List[Trip]:
        wb = self._load()
        ws = self._get_sheet(wb, "Trips", self._init_trips_sheet)
        cols = self._col_map(ws)
        trips: List[Trip] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            try:
                trips.append(self._row_to_trip(row, cols))
            except Exception as exc:
                logger.warning("Skipping malformed trip row: %s", exc)
        return trips

    def get_trip(self, trip_id: str) -> Optional[Trip]:
        for t in self.list_trips():
            if t.trip_id == trip_id:
                return t
        return None

    def save_trip(self, trip: Trip) -> None:
        wb = self._load()
        ws = self._get_sheet(wb, "Trips", self._init_trips_sheet)
        cols = self._col_map(ws)

        # Check for existing row
        for row in ws.iter_rows(min_row=2):
            if row[cols["Trip ID"] - 1].value == trip.trip_id:
                self._write_trip_row(ws, row[0].row, trip, cols)
                self._save(wb)
                return

        # Append new
        self._append_trip_row(ws, trip)
        self._save(wb)

    def delete_trip(self, trip_id: str) -> bool:
        wb = self._load()
        ws = self._get_sheet(wb, "Trips", self._init_trips_sheet)
        cols = self._col_map(ws)
        for row in ws.iter_rows(min_row=2):
            if row[cols["Trip ID"] - 1].value == trip_id:
                ws.delete_rows(row[0].row)
                self._save(wb)
                return True
        return False

    def _row_to_trip(self, row, cols: Dict[str, int]) -> Trip:
        def v(name):
            idx = cols.get(name)
            return row[idx - 1] if idx else None

        return Trip(
            trip_id=str(v("Trip ID")),
            trip_name=str(v("Trip Name")),
            start_date=_safe_date(v("Start Date")),
            end_date=_safe_date(v("End Date")),
            notes=v("Notes"),
            created_at=_safe_datetime(v("Created At")),
            updated_at=_safe_datetime(v("Updated At")),
        )

    def _write_trip_row(self, ws, row_num: int, trip: Trip, cols: Dict[str, int]) -> None:
        data = self._trip_to_row_dict(trip)
        for col_name, value in data.items():
            col_idx = cols.get(col_name)
            if col_idx:
                ws.cell(row=row_num, column=col_idx, value=value)

    def _append_trip_row(self, ws, trip: Trip) -> None:
        data = self._trip_to_row_dict(trip)
        ws.append([data.get(c) for c in TRIP_COLUMNS])

    @staticmethod
    def _trip_to_row_dict(trip: Trip) -> Dict[str, Any]:
        return {
            "Trip ID": trip.trip_id,
            "Trip Name": trip.trip_name,
            "Start Date": trip.start_date.isoformat(),
            "End Date": trip.end_date.isoformat(),
            "Notes": trip.notes,
            "Created At": trip.created_at.isoformat(),
            "Updated At": trip.updated_at.isoformat(),
        }

    # ═══════════════════════════════════════════════════════════════════════════
    # EXPENSES
    # ═══════════════════════════════════════════════════════════════════════════

    def list_expenses(self, f: Optional[ExpenseFilter] = None) -> List[Expense]:
        wb = self._load()
        ws = self._get_sheet(wb, "Expenses", self._init_expenses_sheet)
        cols = self._col_map(ws)
        expenses: List[Expense] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            try:
                exp = self._row_to_expense(row, cols)
                if f is None or self._matches_filter(exp, f):
                    expenses.append(exp)
            except Exception as exc:
                logger.warning("Skipping malformed expense row: %s", exc)
        return sorted(expenses, key=lambda e: e.expense_date, reverse=True)

    def get_expense(self, expense_id: str) -> Optional[Expense]:
        for e in self.list_expenses():
            if e.expense_id == expense_id:
                return e
        return None

    def save_expense(self, expense: Expense) -> None:
        wb = self._load()
        ws = self._get_sheet(wb, "Expenses", self._init_expenses_sheet)
        cols = self._col_map(ws)

        for row in ws.iter_rows(min_row=2):
            if row[cols["Expense ID"] - 1].value == expense.expense_id:
                expense.updated_at = datetime.utcnow()
                self._write_expense_row(ws, row[0].row, expense, cols)
                self._save(wb)
                return

        self._append_expense_row(ws, expense)
        self._save(wb)

    def delete_expense(self, expense_id: str) -> bool:
        wb = self._load()
        ws = self._get_sheet(wb, "Expenses", self._init_expenses_sheet)
        cols = self._col_map(ws)
        for row in ws.iter_rows(min_row=2):
            if row[cols["Expense ID"] - 1].value == expense_id:
                ws.delete_rows(row[0].row)
                self._save(wb)
                return True
        return False

    @staticmethod
    def _matches_filter(exp: Expense, f: ExpenseFilter) -> bool:
        if f.trip_id and exp.trip_id != f.trip_id:
            return False
        if f.date_from and exp.expense_date < f.date_from:
            return False
        if f.date_to and exp.expense_date > f.date_to:
            return False
        if f.expense_type and exp.expense_type != f.expense_type:
            return False
        if f.currency and exp.currency_original != f.currency:
            return False
        if f.submitted is not None and exp.submitted_status != f.submitted:
            return False
        if f.vendor:
            if not exp.vendor_name or f.vendor.lower() not in exp.vendor_name.lower():
                return False
        if f.search:
            q = f.search.lower()
            haystack = " ".join(filter(None, [
                exp.vendor_name,
                exp.receipt_file_name,
                exp.notes,
                str(exp.expense_type.value) if exp.expense_type else "",
            ])).lower()
            if q not in haystack:
                return False
        return True

    def _row_to_expense(self, row, cols: Dict[str, int]) -> Expense:
        def v(name):
            idx = cols.get(name)
            return row[idx - 1] if idx else None

        def safe_enum(cls, val, default):
            if val is None:
                return default
            try:
                return cls(val)
            except ValueError:
                return default

        return Expense(
            expense_id=str(v("Expense ID")),
            trip_id=v("Trip ID"),
            trip_name=v("Trip Name"),
            expense_date=_safe_date(v("Expense Date")),
            vendor_name=v("Vendor Name"),
            expense_type=safe_enum(ExpenseType, v("Expense Type"), ExpenseType.UNKNOWN),
            amount_original=_safe_float(v("Amount Original")) or 0.0,
            currency_original=safe_enum(Currency, v("Currency Original"), Currency.CAD),
            exchange_rate_to_cad=_safe_float(v("Exchange Rate to CAD")),
            amount_cad=_safe_float(v("Amount CAD")) or 0.0,
            gst_original=_safe_float(v("GST Original")),
            gst_cad=_safe_float(v("GST CAD")),
            source_type=safe_enum(SourceType, v("Source Type"), SourceType.MANUAL),
            per_diem_type=safe_enum(PerDiemType, v("Per Diem Type"), None),
            receipt_file_name=v("Receipt File Name"),
            receipt_file_path=v("Receipt File Path"),
            notes=v("Notes"),
            submitted_status=_safe_bool(v("Submitted Status")),
            created_at=_safe_datetime(v("Created At")),
            updated_at=_safe_datetime(v("Updated At")),
        )

    def _write_expense_row(self, ws, row_num: int, exp: Expense, cols: Dict[str, int]) -> None:
        data = self._expense_to_row_dict(exp)
        for col_name, value in data.items():
            col_idx = cols.get(col_name)
            if col_idx:
                ws.cell(row=row_num, column=col_idx, value=value)

    def _append_expense_row(self, ws, exp: Expense) -> None:
        data = self._expense_to_row_dict(exp)
        ws.append([data.get(c) for c in EXPENSE_COLUMNS])

    @staticmethod
    def _expense_to_row_dict(exp: Expense) -> Dict[str, Any]:
        return {
            "Expense ID": exp.expense_id,
            "Trip ID": exp.trip_id,
            "Trip Name": exp.trip_name,
            "Expense Date": exp.expense_date.isoformat(),
            "Vendor Name": exp.vendor_name,
            "Expense Type": exp.expense_type.value if exp.expense_type else None,
            "Amount Original": exp.amount_original,
            "Currency Original": exp.currency_original.value if exp.currency_original else None,
            "Exchange Rate to CAD": exp.exchange_rate_to_cad,
            "Amount CAD": exp.amount_cad,
            "GST Original": exp.gst_original,
            "GST CAD": exp.gst_cad,
            "Source Type": exp.source_type.value if exp.source_type else None,
            "Per Diem Type": exp.per_diem_type.value if exp.per_diem_type else None,
            "Receipt File Name": exp.receipt_file_name,
            "Receipt File Path": exp.receipt_file_path,
            "Notes": exp.notes,
            "Submitted Status": exp.submitted_status,
            "Created At": exp.created_at.isoformat(),
            "Updated At": exp.updated_at.isoformat(),
        }

    # ═══════════════════════════════════════════════════════════════════════════
    # REPORTING
    # ═══════════════════════════════════════════════════════════════════════════

    def summary_by_trip(self) -> List[Dict]:
        expenses = self.list_expenses()
        trips = {t.trip_id: t.trip_name for t in self.list_trips()}
        buckets: Dict[str, Dict] = {}
        for exp in expenses:
            key = exp.trip_id or "__unassigned__"
            name = exp.trip_name or "Unassigned"
            if key not in buckets:
                buckets[key] = {"trip_id": key, "trip_name": name, "total_cad": 0.0, "count": 0}
            buckets[key]["total_cad"] += exp.amount_cad
            buckets[key]["count"] += 1
        return sorted(buckets.values(), key=lambda x: x["total_cad"], reverse=True)

    def summary_by_type(self) -> List[Dict]:
        expenses = self.list_expenses()
        buckets: Dict[str, Dict] = {}
        for exp in expenses:
            key = exp.expense_type.value
            if key not in buckets:
                buckets[key] = {"expense_type": key, "total_cad": 0.0, "count": 0}
            buckets[key]["total_cad"] += exp.amount_cad
            buckets[key]["count"] += 1
        return sorted(buckets.values(), key=lambda x: x["total_cad"], reverse=True)

    def summary_by_month(self) -> List[Dict]:
        expenses = self.list_expenses()
        buckets: Dict[str, Dict] = {}
        for exp in expenses:
            key = exp.expense_date.strftime("%Y-%m")
            if key not in buckets:
                buckets[key] = {"month": key, "total_cad": 0.0, "count": 0}
            buckets[key]["total_cad"] += exp.amount_cad
            buckets[key]["count"] += 1
        return sorted(buckets.values(), key=lambda x: x["month"], reverse=True)

    def total_gst(self) -> float:
        return sum(e.gst_cad or 0.0 for e in self.list_expenses())

    def total_cad(self) -> float:
        return sum(e.amount_cad for e in self.list_expenses())


workbook_service = WorkbookService()
